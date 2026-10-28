"""
Debug script to inspect the structure of .docx files.
Prints paragraph styles, bold runs, and text patterns.
"""

import os
from docx import Document

SOURCE_DIR = os.path.join(os.path.dirname(__file__), '..', 'source')

def inspect_doc(filepath, max_paragraphs=60):
    """Print detailed info about each paragraph in a docx file."""
    doc = Document(filepath)
    print(f"\n{'='*80}")
    print(f"FILE: {os.path.basename(filepath)}")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"{'='*80}")
    
    for i, para in enumerate(doc.paragraphs[:max_paragraphs]):
        text = para.text.strip()
        if not text:
            continue
        
        style_name = para.style.name if para.style else "None"
        has_bold = any(run.bold for run in para.runs if run.text.strip())
        
        # Truncate for readability
        display_text = text[:120] + ('...' if len(text) > 120 else '')
        
        print(f"\n  [P{i}] Style='{style_name}' Bold={has_bold}")
        print(f"       Text: {display_text}")
        
        # Show individual runs for first few paragraphs
        if i < 20:
            for j, run in enumerate(para.runs):
                if run.text.strip():
                    print(f"         Run[{j}]: bold={run.bold} text='{run.text.strip()[:80]}'")

if __name__ == '__main__':
    source_path = os.path.abspath(SOURCE_DIR)
    files = sorted(f for f in os.listdir(source_path) if f.endswith('.docx') and 'Master_Index' not in f)
    
    # Inspect just the first file to understand structure
    if files:
        inspect_doc(os.path.join(source_path, files[0]))
